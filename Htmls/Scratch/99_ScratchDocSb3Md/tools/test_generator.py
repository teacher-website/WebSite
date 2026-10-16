"""用 Playwright 實際操作產生器：上傳 .sb3、切換組別、下載 Word 與 AI 分析包、截圖。

用法：python tools/test_generator.py
輸出：tools/test_out/（截圖、下載的 .docx／.md、summary.json）
"""
import json
import sys
from pathlib import Path

from docx import Document
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
PAGE = ROOT / "貓咪盃程式說明文件產生器.html"
OUT = ROOT / "tools" / "test_out"
PROJ = ROOT.parent  # 01_專案
CASES = [
    ("anim", PROJ / "20260923_森林動物搬家記_溪流上游_Web/student/森林動物搬家記_溪流上游.sb3"),
    ("game", PROJ / "20260923_部落樂舞_火堆旁的節奏_Web/student/部落樂舞_火堆旁的節奏.sb3"),
    ("anim", PROJ / "20260921_Scratch動畫製作教學_Web/student/動畫教學_李昀宸.sb3"),
    ("game", PROJ / "20260921_爆炸頭節奏大師_Web/student/爆炸頭節奏大師_範例.sb3"),
    ("anim", PROJ / "20260924_貓咪甜點店_片頭與選單_Web/student/貓咪甜點店_片頭與選單.sb3"),
    ("game", PROJ / "20260924_手勢辨識_數字與猜拳_Web/student/手勢辨識大挑戰.sb3"),
    # 「其他作品」：不屬於競賽組別的作品
    ("other", PROJ / "20260926_Scratch函式庫_Web/student/sb3/8_遊戲機制.sb3"),
    ("other", PROJ / "20260925_手勢辨識模組教學_Web/student/sb3/模組07_認猜拳.sb3"),
    ("other", PROJ / "20260924_貓咪甜點店_片頭與選單_Web/student/貓咪甜點店_片頭與選單.sb3"),
    ("other", PROJ / "20260923_森林動物搬家記_溪流上游_Web/student/森林動物搬家記_溪流上游.sb3"),
]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    summary, errors = [], []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        for width, label in [(1280, "desktop"), (390, "mobile")]:
            ctx = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
            page = ctx.new_page()
            page.on("console", lambda m: m.type == "error" and errors.append(m.text))
            page.on("pageerror", lambda e: errors.append(str(e)))
            page.goto(PAGE.as_uri())
            cases = CASES if label == "desktop" else CASES[:2] + CASES[6:7]
            for i, (grp, sb3) in enumerate(cases):
                page.check(f'input[name=grp][value={grp}]')
                page.set_input_files("#fileIn", str(sb3))
                page.wait_for_function("n => window.__gen.S.M && window.__gen.S.M.fileName === n && document.querySelector('.paper')", arg=sb3.name, timeout=10000)
                if page.is_visible("#loadErr"):
                    errors.append(f"{sb3.name}: {page.inner_text('#loadErr')}")
                    continue
                info = page.evaluate("""() => { const S = window.__gen.S; return {
                    checks: S.checks.map(c => [c.st, c.title]),
                    doc: JSON.parse(JSON.stringify(S.doc)),
                    aiConds: S.M.aiConds.length, flags: S.M.flags.map(f => f.name), actions: S.M.actions,
                    branches: S.M.branches.length, todo: document.querySelectorAll('.paper .todo').length }; }""")
                tag = f"{label}_{i}_{grp}_{sb3.stem}"
                page.screenshot(path=str(OUT / f"{tag}.png"), full_page=True)
                if label == "desktop":
                    # 在紙本上直接修改一格，確認會寫回資料並出現在 Word 裡
                    first = page.locator(".paper [contenteditable]").first
                    first.click()
                    page.keyboard.press("End")
                    page.keyboard.type("（測試修改）")
                    page.fill('[data-b="name"]', "測試學生")
                    page.wait_for_timeout(700)
                    with page.expect_download() as d:
                        page.click('[data-act="docx"]')
                    docx_path = OUT / f"{tag}.docx"
                    d.value.save_as(docx_path)
                    with page.expect_download() as d:
                        page.click('[data-act="md"]')
                    md_path = OUT / f"{tag}.md"
                    d.value.save_as(md_path)
                    doc = Document(str(docx_path))
                    alltext = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                    info["docx_tables"] = len(doc.tables)
                    info["docx_has_edit"] = "（測試修改）" in alltext
                    info["docx_has_name"] = "測試學生" in alltext
                    if grp == "game":
                        # Word 裡「目標是：」與目標內容要在同一個段落
                        paras = [p.text for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
                        info["docx_goal_inline"] = any(p.startswith("1、 玩家在這個遊戲的目標是：（草稿）") for p in paras)
                    md = md_path.read_text(encoding="utf-8")
                    info["md_kb"] = round(len(md.encode("utf-8")) / 1024, 1)
                    info["md_has_edit"] = "（測試修改）" in md
                    info["md_sections"] = [ln for ln in md.splitlines() if ln.startswith("## ")]
                    if grp == "other":
                        # 作品簡介／操作方式要接在冒號後面；新增角色要點得到、打得進去
                        info["inline_before"] = page.evaluate("""() => [...document.querySelectorAll('.paper td.goal, .paper td.usage')].map(d => {
                            const s = d.querySelector('[contenteditable]');
                            const r = document.createRange(); r.selectNodeContents(d.firstChild);
                            return Math.abs(s.getClientRects()[0].top - r.getClientRects()[0].top) < 6 && s.classList.contains('inl'); })""")
                        page.click('[data-act="addRole"]')
                        page.locator(".paper tr.keep").last.locator("[contenteditable]").last.click()
                        page.keyboard.type("新角色工作")
                        info["new_pair_ok"] = page.evaluate("() => window.__gen.S.doc.roles.at(-1).job") == "新角色工作"
                        info["roles"] = len(info["doc"]["roles"])
                        info["checks_detail"] = page.evaluate("() => window.__gen.S.checks.map(c => c.st + ' ' + c.title + '：' + c.detail)")
                    if grp == "game":
                        # 任務／規則的內容要接在「任務一：」冒號後面（同一行起頭），打字後也要維持
                        inline_js = """() => [document.querySelector('.paper td.goal'), ...document.querySelectorAll('.paper .pair > div')].slice(0, 3).map(d => {
                            const s = d.querySelector('[contenteditable]');
                            const r = document.createRange(); r.selectNodeContents(d.firstChild);  // 冒號前的標題文字
                            return Math.abs(s.getClientRects()[0].top - r.getClientRects()[0].top) < 6 && s.classList.contains('inl'); })"""
                        info["inline_before"] = page.evaluate(inline_js)
                        page.locator(".paper .pair [contenteditable]").first.click()
                        page.keyboard.press("End")
                        page.keyboard.type("X")
                        info["inline_after_typing"] = page.evaluate(inline_js)
                        # 新增一組空白任務，要點得到、打得進去
                        page.click('[data-act="addPair"]')
                        new_rule = page.locator(".paper .pair").last.locator("[contenteditable]").last
                        new_rule.click()
                        page.keyboard.type("新規則")
                        info["new_pair_ok"] = page.evaluate("() => window.__gen.S.doc.pairs.at(-1).rule") == "新規則"
                    # 列印（print 樣式）：用 Chromium 輸出 PDF，確認不是空白、文件內容都在
                    pdf_path = OUT / f"{tag}_print.pdf"
                    page.pdf(path=str(pdf_path), format="A4")
                    import fitz
                    pdf = fitz.open(str(pdf_path))
                    ptext = "".join(pg.get_text() for pg in pdf)
                    info["print_pages"] = pdf.page_count
                    info["print_ok"] = all(k in ptext for k in ("程式說明文件", "基本資料" if grp == "other" else "基本參賽資料", "程式說明", "測試學生")) and ("貓咪盃" in ptext) == (grp != "other") and "上傳作品" not in ptext
                    pdf[0].get_pixmap(dpi=60).save(str(OUT / f"{tag}_print_p1.png"))
                    pdf.close()
                    # 切到另外兩個分頁，確認不會出錯
                    for tab in ("facts", "code", "doc"):
                        page.click(f'[data-tab="{tab}"]')
                info["case"] = tag
                summary.append(info)
            ctx.close()
        browser.close()
    (OUT / "summary.json").write_text(json.dumps({"errors": errors, "cases": summary}, ensure_ascii=False, indent=1), encoding="utf-8")
    for s in summary:
        print(s["case"], "| AI條件", s["aiConds"], "| 旗標", s["flags"], "| 動作", s["actions"], "| 分支", s["branches"], "| 待補", s["todo"],
              "| docx", s.get("docx_tables"), s.get("docx_has_edit"), s.get("docx_has_name"), "| md", s.get("md_kb"), "KB", s.get("md_has_edit"), "| 列印", s.get("print_pages"), "頁", s.get("print_ok"),
              "| 同行", s.get("inline_before"), s.get("inline_after_typing"), "新增", s.get("new_pair_ok"), "| Word 目標同段", s.get("docx_goal_inline"))
        print("   檢核:", " / ".join(f"{st}:{t[:14]}" for st, t in s["checks"]))
    print("console 錯誤：", errors or "無")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
